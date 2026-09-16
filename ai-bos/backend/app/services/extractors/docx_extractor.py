"""
DOCX Word document extractor using python-docx with zipfile fallback.
"""

from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

from app.services.extractors.base import BaseExtractor, ExtractionResult


class DocxExtractor(BaseExtractor):
    def extract(self, file_path: Path) -> ExtractionResult:
        # Try using python-docx first
        try:
            import docx

            doc = docx.Document(str(file_path))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

            # Extract tables if present
            table_texts = []
            for table in doc.tables:
                rows_text = []
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        rows_text.append(" | ".join(cells))
                if rows_text:
                    table_texts.append("\n".join(rows_text))

            combined = []
            if paragraphs:
                combined.append("\n\n".join(paragraphs))
            if table_texts:
                combined.append("\n\n[Tables]\n" + "\n\n".join(table_texts))

            full_text = "\n\n".join(combined)

            return ExtractionResult(
                text=full_text,
                metadata={
                    "paragraph_count": len(paragraphs),
                    "table_count": len(doc.tables),
                    "format": "docx",
                },
            )
        except ImportError:
            # Fallback: parse word/document.xml directly from docx zip package
            try:
                with zipfile.ZipFile(file_path) as z:
                    xml_content = z.read("word/document.xml")
                root = ET.fromstring(xml_content)
                namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                paragraphs = []
                for p in root.findall(".//w:p", namespaces):
                    texts = [node.text for node in p.findall(".//w:t", namespaces) if node.text]
                    if texts:
                        paragraphs.append("".join(texts).strip())

                full_text = "\n\n".join(paragraphs)
                return ExtractionResult(
                    text=full_text,
                    metadata={"paragraph_count": len(paragraphs), "format": "docx_fallback"},
                )
            except Exception as e:
                return ExtractionResult(
                    text="",
                    metadata={"error": str(e), "format": "docx_error"},
                )
